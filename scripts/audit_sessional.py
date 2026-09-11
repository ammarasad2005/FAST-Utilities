#!/usr/bin/env python3
"""
Sessional 1 Transformation Audit
================================
Independently parses each cell of the raw xlsx and cross-checks against
the regular_schedule.json output. Reports any discrepancies.

Bidirectional audit:
1. Forward:  For each raw cell → expected entries → check they exist in JSON
2. Reverse:  For each JSON entry → check it corresponds to a raw cell

Run:
    python3 /home/z/my-project/scripts/audit_sessional.py
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

RAW = "/home/z/my-project/upload/sessional1_fall2026.xlsx"
JSON = "/home/z/my-project/repo/exam-table/public/data/regular_schedule.json"

# ── Config ──────────────────────────────────────────────────────────────────
SHEETS = [("FSC", "FSC", True), ("FSM", "FSM", True), ("FSE", "FSE", False)]
TIME_COLS_FULL = [2, 4, 6, 8, 10, 12, 14, 16]
TIME_COLS_FSE = [2, 4, 6, 8, 10, 12]

# Valid departments (from the parser's DEPARTMENTS list)
VALID_DEPTS = {"CS", "AI", "DS", "CY", "SE", "AF", "FT", "BA", "BBA", "EE", "CE"}

# Grad program prefixes to skip
GRAD_PREFIXES = ("MS", "PhD", "MB", "MSBA", "MEE", "MCI", "MAAIHS", "MYS", "MCS")


def normalize_time(time_str: str) -> str:
    """Same normalization as the transformation script (uses END time's hour)."""
    s = time_str.strip()
    if "AM" in s.upper() or "PM" in s.upper():
        return s
    times = re.findall(r"(\d{1,2}):(\d{2})", s)
    if not times:
        return s
    end_h = int(times[-1][0])
    period = "AM" if 8 <= end_h <= 11 else "PM"
    return f"{s} {period}"


def build_value_map(ws):
    """Build {(row, col): value} with merged cells resolved."""
    vmap = {}
    for row in ws.iter_rows():
        for cell in row:
            vmap[(cell.row, cell.column)] = cell.value
    for merge in ws.merged_cells.ranges:
        tl_val = vmap.get((merge.min_row, merge.min_col))
        for r in range(merge.min_row, merge.max_row + 1):
            for c in range(merge.min_col, merge.max_col + 1):
                if (r, c) != (merge.min_row, merge.min_col):
                    vmap[(r, c)] = tl_val
    return vmap


def find_header_row(ws):
    """Find the row with 'Days' in col A and a time string in col B."""
    for r in range(1, min(ws.max_row + 1, 15)):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        if (a and isinstance(a, str) and "days" in a.lower()
                and b and isinstance(b, str) and re.search(r"\d{1,2}:\d{2}", b)):
            return r
    return None


# ── Independent cell parser (deliberately simple, for cross-checking) ────────

def parse_cell_independently(cell_text: str, school: str, date_str: str, time_str: str):
    """Parse a cell's text and return a list of expected entry tuples.

    Each entry is: (date, time, courseCode, courseName, batch, department, school)

    This parser is INDEPENDENT from the transformation script — it uses a
    different approach (line-by-line manual parsing) to cross-check.
    """
    text = cell_text.strip()
    if not text or "reserved" in text.lower():
        return []

    lines = [l.rstrip() for l in text.split("\n")]
    lines = [l for l in lines if l.strip()]  # remove blank lines
    if not lines:
        return []

    # ── Step 1: Extract course code from first line ──
    first_line = lines[0].strip()
    code_match = re.match(r"^([A-Za-z]{2,4}[\s\-]?\d{4})", first_line)
    if not code_match:
        return []  # No course code found — skip
    code_raw = code_match.group(1)
    course_code = re.sub(r"[\s\-]", "", code_raw).upper()

    # ── Step 2: Extract course name ──
    # The name is everything after the code on the first line, possibly
    # continuing to the second line if the first line is just the code.
    rest_of_first = first_line[len(code_raw):].strip()
    rest_of_first = re.sub(r"^[\s\-]+", "", rest_of_first).strip()  # strip leading " - "

    if rest_of_first:
        # Check if rest_of_first contains a program pattern — if so, split it
        prog_in_first = re.search(
            r"(?:BS\s*\(|\bBBA\b|\bBCE\b|\bBEE\b|\bMS\s*\(|\bPhD\b|\bMB[\s-]|\bMSBA[\s-]|\bMEE[\s-])",
            rest_of_first,
        )
        if prog_in_first:
            course_name = rest_of_first[:prog_in_first.start()].strip().strip("-").strip()
            # The program part + rest of first line
            remaining_text = rest_of_first[prog_in_first.start():] + "\n" + "\n".join(lines[1:])
        else:
            # No program in first line — name is the rest (minus trailing batch year)
            course_name = re.sub(r"\s+\d{4}$", "", rest_of_first).strip()
            remaining_text = "\n".join(lines[1:])
    else:
        # First line was just the code — name is on the second line
        if len(lines) > 1:
            second = lines[1].strip()
            course_name = re.sub(r"\s+\d{4}$", "", second).strip()
            remaining_text = "\n".join(lines[2:])
        else:
            return []  # No name found

    # ── Step 3: Extract batch year ──
    # Look for a standalone 4-digit year on the last few lines, or any 4-digit year
    batch = ""
    remaining_lines = remaining_text.strip().split("\n") if remaining_text.strip() else []
    for line in reversed(remaining_lines[-3:]):
        line = line.strip()
        if re.match(r"^\d{4}$", line):
            batch = line
            break
    if not batch:
        m = re.search(r"\b(20\d{2})\b", text)
        if m:
            batch = m.group(1)

    if not batch:
        return []  # No batch year → skip (matches transformation script behavior)

    # ── Step 4: Parse BS programs ──
    # Collect all program lines from remaining_text + any program part from first line
    full_program_text = remaining_text

    programs = []  # list of (department, sections)
    seen = set()

    def add_dept(dept, sections):
        dept = dept.upper().strip()
        sections = (sections or "").strip().rstrip(",").strip()
        key = (dept, sections)
        if dept and dept in VALID_DEPTS and key not in seen:
            seen.add(key)
            programs.append((dept, sections))

    # BS(dept) (sections) — standard format
    # NOTE: [^)\n]* to prevent greedy matching across newlines (same fix as
    # the transformation script)
    for m in re.finditer(r"BS\s*\(\s*(CS|AI|DS|CY|SE|AF|FT|BA|EE|CE)\s*\)\s*(?:\(([^)\n]*)\))?", full_program_text, re.IGNORECASE):
        add_dept(m.group(1), m.group(2))

    # BBA standalone
    for m in re.finditer(r"\bBBA\b\s*[-]?\s*(?:\(([^)]*)\)|([A-Za-z0-9,]+))?", full_program_text, re.IGNORECASE):
        add_dept("BBA", m.group(1) or m.group(2))

    # BCE → CE
    for m in re.finditer(r"\bBCE\b\s*[-]?\s*(?:\(([^)]*)\)|([A-Za-z0-9,]+))?", full_program_text, re.IGNORECASE):
        add_dept("CE", m.group(1) or m.group(2))

    # BEE → EE
    for m in re.finditer(r"\bBEE\b\s*[-]?\s*(?:\(([^)]*)\)|([A-Za-z0-9,]+))?", full_program_text, re.IGNORECASE):
        add_dept("EE", m.group(1) or m.group(2))

    # Also handle inline BS programs with sections after a space or dash
    # (e.g., "BS(AI) A,B,C" or "BS(AF)-A,B")
    for m in re.finditer(r"BS\s*\(\s*(CS|AI|DS|CY|SE|AF|FT|BA|EE|CE)\s*\)\s*[-]?\s*([A-Za-z0-9,]+)", full_program_text, re.IGNORECASE):
        add_dept(m.group(1), m.group(2))

    if not programs:
        return []  # No BS programs found (maybe all grad) → skip

    # ── Step 5: Build expected entries ──
    entries = []
    for dept, sections in programs:
        entries.append({
            "date": date_str,
            "time": time_str,
            "courseCode": course_code,
            "courseName": course_name,
            "batch": batch,
            "department": dept,
            "school": school,
        })
    return entries


def main():
    print("=" * 80)
    print("SESSIONAL 1 AUDIT — cell-by-cell cross-check")
    print("=" * 80)

    # Load raw xlsx
    wb = load_workbook(RAW, data_only=True)
    print(f"\n📄 Loaded raw xlsx: {RAW}")
    print(f"   Sheets: {wb.sheetnames}")

    # Load JSON output
    with open(JSON) as f:
        json_entries = json.load(f)
    print(f"📊 Loaded JSON: {len(json_entries)} entries")

    # Build a lookup set of JSON entries for fast comparison
    # Key: (date, time, courseCode, batch, department, school)
    # Value: list of entries (for checking courseName match)
    json_lookup = defaultdict(list)
    for e in json_entries:
        key = (e["date"], e["time"], e["courseCode"], e["batch"], e["department"], e["school"])
        json_lookup[key].append(e)

    # ── FORWARD AUDIT: raw cells → expected entries → check in JSON ──
    expected_entries = []
    missing_entries = []
    name_mismatches = []
    cells_audited = 0
    cells_with_data = 0
    cells_skipped_no_batch = 0
    cells_skipped_no_programs = 0
    cells_skipped_grad_only = 0

    for sheet_name, school, has_late_slots in SHEETS:
        ws = wb[sheet_name]
        vmap = build_value_map(ws)
        header_row = find_header_row(ws)
        if not header_row:
            print(f"\n⚠️  No header row in {sheet_name}")
            continue

        time_cols = TIME_COLS_FULL if has_late_slots else TIME_COLS_FSE
        time_slots = {}
        for col in time_cols:
            val = vmap.get((header_row, col))
            if val and isinstance(val, str) and val.strip():
                time_slots[col] = normalize_time(val.strip())

        print(f"\n{'─'*60}")
        print(f"SHEET: {sheet_name} (school={school}, header_row={header_row})")
        print(f"  Time slots: {time_slots}")

        current_date = None
        for r in range(header_row + 1, ws.max_row + 1):
            cell_a = vmap.get((r, 1))

            if isinstance(cell_a, dt.datetime):
                current_date = cell_a
            elif isinstance(cell_a, str) and cell_a.strip().lower().startswith("note:"):
                break
            elif cell_a is None or (isinstance(cell_a, str) and not cell_a.strip()):
                has_data = any(vmap.get((r, col)) for col in time_cols)
                if not has_data:
                    continue

            if current_date is None:
                continue

            date_str = current_date.strftime("%d/%m/%Y")

            for col in time_cols:
                time_str = time_slots.get(col)
                if not time_str:
                    continue

                cell_val = vmap.get((r, col))
                if not cell_val or not str(cell_val).strip():
                    continue

                cells_audited += 1
                cell_text = str(cell_val).strip()

                if "reserved" in cell_text.lower():
                    continue

                cells_with_data += 1
                expected = parse_cell_independently(cell_text, school, date_str, time_str)

                if not expected:
                    # Check if it's because no batch or no programs
                    batch = re.search(r"\b20\d{2}\b", cell_text)
                    has_bs = re.search(r"BS\s*\(|\bBBA\b|\bBCE\b|\bBEE\b", cell_text, re.IGNORECASE)
                    if not batch:
                        cells_skipped_no_batch += 1
                        reason = "no batch year"
                    elif not has_bs:
                        cells_skipped_grad_only += 1
                        reason = "grad-only (no BS programs)"
                    else:
                        cells_skipped_no_programs += 1
                        reason = "BS programs found but not parsed"
                    print(f"  ⚠️  SKIP ({reason}): {sheet_name} row {r} col {col} — {cell_text[:80]!r}")
                    continue

                for entry in expected:
                    expected_entries.append(entry)
                    key = (entry["date"], entry["time"], entry["courseCode"], entry["batch"], entry["department"], entry["school"])
                    if key not in json_lookup:
                        missing_entries.append((sheet_name, r, col, entry))
                    else:
                        # Check courseName match
                        json_matches = json_lookup[key]
                        for jm in json_matches:
                            if jm["courseName"] != entry["courseName"]:
                                name_mismatches.append({
                                    "cell": f"{sheet_name} row {r} col {col}",
                                    "expected_name": entry["courseName"],
                                    "actual_name": jm["courseName"],
                                    "entry": entry,
                                })

    # ── REVERSE AUDIT: JSON entries → check they correspond to a raw cell ──
    expected_keys = set()
    for e in expected_entries:
        key = (e["date"], e["time"], e["courseCode"], e["batch"], e["department"], e["school"])
        expected_keys.add(key)

    extra_entries = []
    for e in json_entries:
        key = (e["date"], e["time"], e["courseCode"], e["batch"], e["department"], e["school"])
        if key not in expected_keys:
            extra_entries.append(e)

    # ── REPORT ──
    print(f"\n{'='*80}")
    print("AUDIT REPORT")
    print(f"{'='*80}")
    print(f"\n📊 Cell statistics:")
    print(f"  Cells with data:           {cells_with_data}")
    print(f"  Cells skipped (no batch):  {cells_skipped_no_batch}")
    print(f"  Cells skipped (grad-only): {cells_skipped_grad_only}")
    print(f"  Cells skipped (no parse):  {cells_skipped_no_programs}")
    print(f"\n📋 Entry statistics:")
    print(f"  Expected entries (from raw): {len(expected_entries)}")
    print(f"  Actual entries (in JSON):    {len(json_entries)}")
    print(f"\n🔍 Discrepancies:")

    if missing_entries:
        print(f"\n  ❌ MISSING ({len(missing_entries)} entries expected but not in JSON):")
        for sheet, r, col, entry in missing_entries[:20]:
            print(f"     {sheet} R{r}C{col}: {entry['date']} {entry['time']} {entry['courseCode']} {entry['courseName'][:30]} batch={entry['batch']} dept={entry['department']}")
    else:
        print(f"\n  ✅ No missing entries — all expected entries found in JSON")

    if extra_entries:
        print(f"\n  ❌ EXTRA ({len(extra_entries)} entries in JSON but not expected from raw):")
        for e in extra_entries[:20]:
            print(f"     {e['date']} {e['time']} {e['courseCode']} {e['courseName'][:30]} batch={e['batch']} dept={e['department']}")
    else:
        print(f"\n  ✅ No extra entries — all JSON entries correspond to a raw cell")

    if name_mismatches:
        print(f"\n  ❌ NAME MISMATCHES ({len(name_mismatches)} entries with different courseName):")
        for nm in name_mismatches[:10]:
            print(f"     {nm['cell']}: expected={nm['expected_name']!r} actual={nm['actual_name']!r}")
    else:
        print(f"\n  ✅ No name mismatches — all courseNames match")

    # ── FINAL VERDICT ──
    total_discrepancies = len(missing_entries) + len(extra_entries) + len(name_mismatches)
    print(f"\n{'='*80}")
    if total_discrepancies == 0:
        print(f"✅ AUDIT PASSED — 0 discrepancies. All {len(expected_entries)} expected entries match the JSON.")
    else:
        print(f"❌ AUDIT FAILED — {total_discrepancies} discrepancy(ies) found:")
        print(f"   Missing: {len(missing_entries)}")
        print(f"   Extra:   {len(extra_entries)}")
        print(f"   Name:    {len(name_mismatches)}")
    print(f"{'='*80}")
    return 0 if total_discrepancies == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
